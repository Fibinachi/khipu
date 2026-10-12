#!/usr/bin/env python3
"""
Parse KFG Excel files into kfg-khipu.db, mirroring okr-khipu.db structure.
KHIPU_ID scheme: 2000000 + KH number (e.g. KH0002 -> 2000002).
"""
import os
import re
import sqlite3
import sys

import openpyxl

RAW_DIR = os.path.expanduser("~/workspace/quipu-analysis/kfg_raw")
DB_PATH = os.path.expanduser("~/workspace/quipu-analysis/kfg-khipu.db")

SCHEMA = """
CREATE TABLE IF NOT EXISTS khipu_main (
  KHIPU_ID INTEGER PRIMARY KEY,
  KH_NUM TEXT,
  ALIASES TEXT,
  CONTRIBUTORS TEXT,
  PROVENANCE TEXT,
  REGION TEXT,
  MUSEUM_NAME TEXT,
  MUSEUM_NUM TEXT,
  MUSEUM_DESCR TEXT,
  NICKNAME TEXT,
  INVESTIGATOR_NUM TEXT,
  NOTES TEXT,
  KFG_URL TEXT
);
CREATE TABLE IF NOT EXISTS primary_cord (
  KHIPU_ID INTEGER PRIMARY KEY,
  STRUCTURE TEXT,
  THICKNESS TEXT,
  LENGTH TEXT,
  COLOR TEXT,
  FIBER TEXT,
  BEGINNING TEXT,
  TERMINATION TEXT,
  TWIST TEXT,
  NOTES TEXT
);
CREATE TABLE IF NOT EXISTS cord (
  KHIPU_ID INTEGER,
  CORD_ID INTEGER PRIMARY KEY,
  CORD_NAME TEXT,
  CORD_LEVEL INTEGER,
  PENDANT_FROM INTEGER,
  TWIST TEXT,
  ATTACHMENT_TYPE TEXT,
  TERMINATION TEXT,
  THICKNESS TEXT,
  CORD_LENGTH TEXT,
  CORD_ORDINAL INTEGER,
  KNOT_VALUE TEXT,
  ALT_VALUE TEXT,
  POSITION TEXT,
  CORD_NOTES TEXT
);
CREATE INDEX IF NOT EXISTS idx_kfg_cord_khipu ON cord(KHIPU_ID);
CREATE INDEX IF NOT EXISTS idx_kfg_cord_pendant ON cord(PENDANT_FROM);
CREATE TABLE IF NOT EXISTS cord_color (
  KHIPU_ID INTEGER,
  CORD_ID INTEGER,
  COLOR_SPEC TEXT,
  COLOR_CD_1 TEXT
);
CREATE TABLE IF NOT EXISTS knot (
  KHIPU_ID INTEGER,
  CORD_ID INTEGER,
  KNOT_ORDINAL INTEGER,
  TYPE_CODE TEXT,
  NUM_TURNS TEXT,
  POSITION TEXT,
  TWIST_DIR TEXT,
  KNOT_VALUE TEXT
);
CREATE TABLE IF NOT EXISTS kh_crosswalk (
  KH_NUM TEXT PRIMARY KEY,
  OKR_KHIPU_ID INTEGER,
  KFG_KHIPU_ID INTEGER
);
"""

def parse_nv_sheet(ws):
    """Parse Name:Value single-column sheet into dict."""
    d = {}
    for row in ws.iter_rows(values_only=True):
        if not row[0]:
            continue
        cell = str(row[0]).strip()
        if cell.startswith("!--"):
            continue
        if ":" in cell:
            k, v = cell.split(":", 1)
            d[k.strip()] = v.strip()
    return d

def parse_cord_name(name):
    """'p1s2s3' -> (level, parent_name). 'p1' -> (1, None)."""
    parts = name.split("s")
    # parts[0] = 'p1', rest are subsidiary indices
    level = len(parts)
    if level == 1:
        return 1, None
    parent = "s".join(parts[:-1])
    return level, parent

KNOT_RE = re.compile(r"(\d*)([A-Z]+)\(([0-9.]+),([A-Z?]+)\)(?:,([0-9]+))?(?:\[([A-Z]+)\])?")

def parse_knots(knot_str):
    """'4L(10.5,U),4;3S(7.5,U),30' -> list of dicts."""
    if not knot_str:
        return []
    out = []
    for i, chunk in enumerate(str(knot_str).split(";")):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = KNOT_RE.match(chunk)
        if m:
            turns, ktype, pos, twist, val, annot = m.groups()
            out.append({"ordinal": i + 1, "turns": turns or "", "type": ktype,
                        "pos": pos, "twist": twist, "value": val or "",
                        "annot": annot or ""})
        else:
            out.append({"ordinal": i + 1, "turns": "", "type": "UNPARSED",
                        "pos": "", "twist": "", "value": "", "raw": chunk})
    return out

def color_cd1(spec):
    """Extract primary color code from spec like 'W(0.0-25.0)/MB(25.0-41.0)' or 'W:MB'."""
    if not spec:
        return ""
    # Take first segment before / then strip operators and parens
    seg = str(spec).split("/")[0]
    seg = re.split(r"[:\-_%^]", seg)[0]
    seg = re.sub(r"\(.*", "", seg).strip()
    return seg

def process_file(path, kh_num, conn):
    khipu_id = 2000000 + int(kh_num[2:])
    wb = openpyxl.load_workbook(path, data_only=True)
    cur = conn.cursor()

    # Khipu sheet
    k = parse_nv_sheet(wb["Khipu"])
    cur.execute(
        """INSERT OR REPLACE INTO khipu_main
           (KHIPU_ID, KH_NUM, ALIASES, CONTRIBUTORS, PROVENANCE, REGION,
            MUSEUM_NAME, MUSEUM_NUM, MUSEUM_DESCR, NICKNAME, INVESTIGATOR_NUM, KFG_URL)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
        (khipu_id, kh_num, k.get("Aliases", ""), k.get("Contributors", ""),
         k.get("Provenance", ""), k.get("Region", ""), k.get("Museum Name", ""),
         k.get("Museum Number", ""), k.get("Museum Description", ""),
         k.get("Name", ""), k.get("Aliases", ""), k.get("KFG URL", "")))

    # PrimaryCord sheet
    if "PrimaryCord" in wb.sheetnames:
        p = parse_nv_sheet(wb["PrimaryCord"])
        cur.execute(
            """INSERT OR REPLACE INTO primary_cord
               (KHIPU_ID, STRUCTURE, THICKNESS, LENGTH, COLOR, FIBER,
                BEGINNING, TERMINATION, TWIST, NOTES)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (khipu_id, p.get("Structure", ""), p.get("Thickness", ""),
             p.get("Length", ""), p.get("Color", ""), p.get("Fiber", ""),
             p.get("Beginning", ""), p.get("Termination", ""),
             p.get("Twist", ""), p.get("Notes", "")))

    # Cords sheet
    if "Cords" not in wb.sheetnames:
        return khipu_id, 0
    ws = wb["Cords"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    # First pass: assign CORD_IDs
    cord_ids = {}  # cord_name -> cord_id
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[idx["Cord_Name"]]:
            continue
        rows.append(r)
    # Primary cord gets CORD_ID = khipu_id * 100000 (namespace separator)
    pcord_id = khipu_id * 100000
    for j, r in enumerate(rows):
        cname = str(r[idx["Cord_Name"]]).strip()
        cord_ids[cname] = khipu_id * 100000 + j + 1

    n_cords = 0
    for j, r in enumerate(rows):
        cname = str(r[idx["Cord_Name"]]).strip()
        cid = cord_ids[cname]
        level, parent = parse_cord_name(cname)
        if parent is None:
            pendant_from = pcord_id
        else:
            pendant_from = cord_ids.get(parent)

        def g(col):
            v = r[idx[col]] if col in idx else None
            return "" if v is None else str(v).strip()

        cur.execute(
            """INSERT OR REPLACE INTO cord
               (KHIPU_ID, CORD_ID, CORD_NAME, CORD_LEVEL, PENDANT_FROM, TWIST,
                ATTACHMENT_TYPE, TERMINATION, THICKNESS, CORD_LENGTH, CORD_ORDINAL,
                KNOT_VALUE, ALT_VALUE, POSITION, CORD_NOTES)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (khipu_id, cid, cname, level, pendant_from, g("Twist"),
             g("Attachment"), g("Termination"), g("Thickness"), g("Length"),
             j + 1, g("Value"), g("Alt_Value"), g("Position"), g("Notes")))

        # Color
        cspec = g("Color")
        cur.execute(
            "INSERT INTO cord_color (KHIPU_ID, CORD_ID, COLOR_SPEC, COLOR_CD_1) VALUES (?,?,?,?)",
            (khipu_id, cid, cspec, color_cd1(cspec)))

        # Knots
        for kn in parse_knots(g("Knots")):
            cur.execute(
                """INSERT INTO knot
                   (KHIPU_ID, CORD_ID, KNOT_ORDINAL, TYPE_CODE, NUM_TURNS,
                    POSITION, TWIST_DIR, KNOT_VALUE)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (khipu_id, cid, kn["ordinal"], kn["type"], kn["turns"],
                 kn["pos"], kn["twist"], kn["value"]))
        n_cords += 1

    return khipu_id, n_cords

def main():
    if os.path.exists(DB_PATH):
        os.remove(DB_PATH)
    conn = sqlite3.connect(DB_PATH)
    conn.executescript(SCHEMA)

    # Build OKR crosswalk from local DB
    okr = sqlite3.connect(os.path.expanduser("~/workspace/quipu-analysis/okr-khipu.db"))
    okr_map = dict(okr.execute(
        "SELECT OKR_NUM, KHIPU_ID FROM khipu_main WHERE OKR_NUM IS NOT NULL AND OKR_NUM != ''").fetchall())
    okr.close()

    files = sorted(f for f in os.listdir(RAW_DIR) if f.endswith(".xlsx"))
    total_cords = 0
    for f in files:
        kh_num = f.replace(".xlsx", "")
        try:
            khipu_id, n = process_file(os.path.join(RAW_DIR, f), kh_num, conn)
            total_cords += n
            conn.execute(
                "INSERT OR REPLACE INTO kh_crosswalk (KH_NUM, OKR_KHIPU_ID, KFG_KHIPU_ID) VALUES (?,?,?)",
                (kh_num, okr_map.get(kh_num), khipu_id))
            print(f"  {kh_num}: {n} cords", flush=True)
        except Exception as e:
            print(f"  ERROR {kh_num}: {e}", flush=True)
    conn.commit()

    # Also add crosswalk entries for OKR-only khipu
    for okr_num, okr_id in okr_map.items():
        conn.execute(
            "INSERT OR IGNORE INTO kh_crosswalk (KH_NUM, OKR_KHIPU_ID) VALUES (?,?)",
            (okr_num, okr_id))
    conn.commit()

    nk = conn.execute("SELECT COUNT(*) FROM khipu_main").fetchone()[0]
    nc = conn.execute("SELECT COUNT(*) FROM cord").fetchone()[0]
    nkn = conn.execute("SELECT COUNT(*) FROM knot").fetchone()[0]
    print(f"\nParsed {nk} khipu, {nc} cords, {nkn} knots into {DB_PATH}")
    conn.close()

if __name__ == "__main__":
    main()
